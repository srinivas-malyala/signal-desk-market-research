from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "Signal_Desk_Capstone_Evidence.docx"
SCREENSHOTS = ROOT / "submission" / "screenshots"
ARCHITECTURE = ROOT / "proposal" / "signal-desk-capstone-architecture.png"

INK = "17252B"
TEAL = "0B7A68"
PALE_TEAL = "EAF5F2"
PALE_GRAY = "F4F5F4"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def style_table(table, widths: list[float], header_fill: str = INK) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cell, header_fill)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Aptos"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
                    run.font.size = Pt(9)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if c_idx == 0 and r_idx > 0:
                cell.paragraphs[0].runs[0].font.bold = True


def set_image_alt_text(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(70, 76, 79)


def add_picture(doc: Document, path: Path, width: float, alt: str, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_image_alt_text(shape, alt)
    add_caption(doc, caption)


def add_link(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(r_pr)
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        p_bdr.append(node)
    p_pr.append(p_bdr)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13

    for style_name, size, before, after in (
        ("Title", 30, 0, 18),
        ("Subtitle", 13, 0, 12),
        ("Heading 1", 19, 18, 8),
        ("Heading 2", 13.5, 12, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Subtitle" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if style_name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True

    title_p_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    doc.styles["Caption"].font.name = "Aptos"
    doc.styles["Caption"].font.color.rgb = RGBColor(70, 76, 79)

    for sec in doc.sections:
        add_page_number(sec.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("DATABRICKS AI CAPSTONE SUBMISSION")
    r.font.name = "Aptos"
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r.font.letter_spacing = Pt(1.2) if hasattr(r.font, "letter_spacing") else None

    title = doc.add_paragraph(style="Title")
    title.add_run("Signal Desk Capstone Evidence Report")
    remove_paragraph_borders(title)

    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("Functionality deployment architecture and rubric evidence")

    meta = doc.add_table(rows=5, cols=2)
    meta_data = [
        ("Author", "Srinivas Malyala"),
        ("Application", "Signal Desk AI Stock Market Research Assistant"),
        ("Verified user", "malyalasrinivas@gmail.com"),
        ("Live application", "signal-desk-frontend-s88i.onrender.com"),
        ("Evidence date", "October 7, 2026"),
    ]
    for i, (label, value) in enumerate(meta_data):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
    style_table(meta, [1.45, 5.15], header_fill=INK)
    for cell in meta.rows[0].cells:
        set_cell_shading(cell, PALE_TEAL)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string(INK)
    for row in meta.rows:
        set_cell_shading(row.cells[0], PALE_TEAL)
        row.cells[0].paragraphs[0].runs[0].font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Submission purpose. ").bold = True
    p.add_run(
        "This report gives graders direct evidence for every rubric category and shows how the working application, data platform, agent actions, analytics, deployment, and Big Data characteristics fit together."
    )

    p = doc.add_paragraph()
    p.add_run("Primary conclusion. ").bold = True
    p.add_run(
        "The available code, accepted workspace results, live deployment, and current automated verification support the top rubric band in all eight categories. The report distinguishes live observations from repository acceptance evidence."
    )

    p = doc.add_paragraph()
    p.add_run("Live link. ").bold = True
    add_link(p, "Open Signal Desk", "https://signal-desk-frontend-s88i.onrender.com/")
    doc.add_page_break()


def add_overview(doc: Document) -> None:
    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph(
        "Signal Desk is an authenticated stock research application that combines live market data, SEC and article evidence, an action-taking MCP agent, user-owned Lakebase state, and Databricks analytics. A researcher can inspect price performance, compare peers, retrieve attributable filing and news passages, maintain a watchlist, save research, and review operational analytics from one interface."
    )

    facts = doc.add_table(rows=1, cols=3)
    for i, value in enumerate(("Evidence area", "Demonstrated result", "Primary proof")):
        facts.cell(0, i).text = value
    rows = [
        ("Scale", "1,255,489 unique Silver market rows", "Accepted market pipeline reconciliation"),
        ("Variety", "Structured, semi-structured, and unstructured data", "OHLCV, XBRL, JSON, filings, and news"),
        ("Agent", "Six retrieval and three confirmed write tools", "MCP contracts and 10 of 10 Supervisor evaluation"),
        ("Analytics", "Five Gold metric families", "Lakebase history through Bronze and Silver to Gold"),
        ("Quality", "281 tests and Ruff pass", "Local verification on October 7, 2026"),
        ("Deployment", "Frontend and MCP live on Render", "Live application and Render deployment evidence"),
    ]
    for values in rows:
        cells = facts.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(facts, [1.35, 2.75, 2.55])

    doc.add_page_break()
    doc.add_heading("Architecture", level=1)
    doc.add_paragraph(
        "Render hosts the Flask frontend and FastMCP application tier. Databricks owns the Spark pipelines, Delta tables, SQL analytics, AI Search, and Agent Bricks Supervisor. Lakebase provides operational storage and market-serving data. Short-lived signed assertions carry authenticated browser identity to MCP; a fixed Supervisor credential protects agent calls, and the owner-approved shared workspace identity supports bounded AI Search and Gold reads."
    )
    add_picture(
        doc,
        ARCHITECTURE,
        6.7,
        "Signal Desk architecture showing browser, Render services, Databricks processing, Massive data, and Lakebase storage.",
        "Figure 1  Signal Desk split-host architecture",
    )


def add_score_map(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Rubric Evidence Summary", level=1)
    doc.add_paragraph(
        "The table below identifies the highest rubric band supported by the submitted evidence. It is an evidence map, not a guaranteed grade. Detailed implementation and verification references follow."
    )
    table = doc.add_table(rows=1, cols=4)
    headers = ("Category", "Points", "Top-band evidence", "Where to verify")
    for i, value in enumerate(headers):
        table.cell(0, i).text = value
    evidence = [
        (
            "Spark data pipeline",
            "15",
            "Reusable Spark ingestion, validation, enrichment, quarantine, and Gold outputs",
            "pipelines and resources",
        ),
        (
            "Third-party API",
            "10",
            "Massive and SEC integrations with secrets, retries, caching, validation, and shared quota controls",
            "clients, ingestion, migration 0005",
        ),
        (
            "Lakebase model",
            "15",
            "Normalized operational model with keys, constraints, indexes, ownership, audit fields, and idempotency",
            "migrations and lakebase modules",
        ),
        (
            "Action-taking agent",
            "20",
            "Grounded retrieval plus three protected writes; deployed Supervisor passes 10 of 10 cases",
            "MCP server, prompt, evaluations",
        ),
        (
            "Analytics pipeline",
            "10",
            "Incremental history processing into privacy-safe Silver and five useful Gold metric families",
            "analytics pipeline and acceptance tool",
        ),
        (
            "Frontend workflow",
            "10",
            "Authenticated end-to-end research, watchlist, memory, evidence links, errors, and analytics",
            "live screenshots and dashboard code",
        ),
        (
            "Deployed application",
            "5",
            "Live Render frontend and MCP with OIDC, health checks, Blueprint configuration, and runbook",
            "live URL, render.yaml, runbook",
        ),
        (
            "Two Big Data Vs",
            "15",
            "Volume above one million rows and meaningful unstructured document processing",
            "coverage results and research index",
        ),
    ]
    for values in evidence:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(table, [1.45, 0.55, 3.4, 1.25])
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=60, start=95, bottom=60, end=95)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(8.3)
    for row in table.rows[1:]:
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_category(doc: Document, title: str, points: str, body: str, proof: list[str], results: list[str]) -> None:
    doc.add_page_break()
    doc.add_heading(f"{title}  {points}", level=1)
    doc.add_paragraph(body)
    doc.add_heading("Implementation Evidence", level=2)
    for item in proof:
        add_bullet(doc, item)
    doc.add_heading("Demonstrated Result", level=2)
    for item in results:
        add_bullet(doc, item)


def add_detailed_evidence(doc: Document) -> None:
    add_category(
        doc,
        "Spark Data Pipeline",
        "15 points",
        "Separate Lakeflow Spark Declarative Pipelines process market, SEC, article, document, and activity data. The transformations use explicit schemas, deterministic keys, deduplication, quarantine paths, reconciliation, and reusable Bronze, Silver, and Gold layers.",
        [
            "Market processing: pipelines/market_bars_classified.py, silver_market_bars.py, silver_market_quarantine.py, and gold_market_data_coverage.py.",
            "Research processing: bronze_sec_company_facts.py, bronze_sec_filing_documents.py, silver_sec_facts.py, silver_research_articles.py, and silver_research_chunks.py.",
            "Deployment definitions: resources/market_pipeline.pipeline.yml and resources/research_pipeline.pipeline.yml.",
        ],
        [
            "1,255,677 Bronze market rows reconciled to 1,255,489 unique Silver rows plus 188 deterministic quarantines across 81 manifest dates.",
            "The bounded research run produced 2 companies, 12 filings, 57,806 facts, 86 articles, 549 article-ticker links, and 507 traceable chunks with zero integrity violations.",
            "Cached research ingestion reran without external API calls, demonstrating a repeatable workflow.",
        ],
    )

    add_category(
        doc,
        "Third Party API Integration",
        "10 points",
        "Massive Stocks supplies company, price, fundamentals, and news data. SEC EDGAR supplies company submissions, filing documents, and structured facts. Keys and identifying contacts remain in secrets or runtime configuration rather than source control.",
        [
            "mcp_server/massive_client.py implements bounded calls, response validation, retries, plan-aware errors, and safe result envelopes.",
            "ingestion/market_backfill.py, ingestion/sec_client.py, and ingestion/article_landing.py provide immutable landing, checkpoints, and cached reruns.",
            "migrations/0005_massive_rate_limit.sql and the shared Lakebase ledger coordinate rate limits across Render and Databricks jobs.",
        ],
        [
            "A cross-host acceptance run recorded four Render attempts and one job attempt, never exceeded four acquisitions in a rolling minute, and delayed attempt five by 60.225 seconds.",
            "Malformed or unavailable source data is labeled safely; plan-restricted values are never converted to zero or fabricated.",
        ],
    )

    add_category(
        doc,
        "Lakebase Data Model",
        "15 points",
        "The shared PostgreSQL schema uses suffix-isolated operational tables for users, watchlists, companies, price snapshots, articles, notes, reports, agent events, traces, quota attempts, and serving copies. Versioned migrations are idempotent and checksum protected.",
        [
            "mcp_server/migrations/0001_operational_core.sql through 0005_massive_rate_limit.sql define the model, CDC settings, indexes, and operational fields.",
            "mcp_server/lakebase.py and mcp_server/lakebase_serving.py provide bounded pools, ownership-scoped queries, and serving access.",
            "mcp_server/action_service.py enforces validation, confirmation, idempotency, and changed-payload rejection.",
        ],
        [
            "PostgreSQL 17 or newer connectivity, 16 owned tables, five migrations, and six full-replica-identity tables were verified.",
            "Two real Google principals proved isolated watchlists, notes, reports, and traces through the deployed UI; disposable state was then removed exactly.",
        ],
    )

    add_category(
        doc,
        "Action Taking AI Agent",
        "20 points",
        "FastMCP exposes six retrieval tools and three protected write tools. The Agent Bricks Supervisor selects tools, validates requests, links evidence, asks for confirmation before consequential actions, and reports bounded failures without inventing results.",
        [
            "Retrieval tools cover price performance, company research, peer comparison, watchlists, semantic evidence, and notable updates.",
            "Write tools add or remove watchlist items and save research notes or analysis reports in Lakebase.",
            "agent/system_prompt.md defines source selection, ambiguity handling, confirmation, idempotency, and investment-advice boundaries.",
        ],
        [
            "Exact replays return the original action result; the same idempotency key with a changed payload is rejected.",
            "The deployed governed Supervisor passed all 10 live evaluation cases and exposes the exact nine-tool MCP contract.",
            "Every research card includes a source type, date, title, and link for user inspection.",
        ],
    )

    add_category(
        doc,
        "Analytics Pipeline",
        "10 points",
        "Lakebase history tables feed an incremental Databricks analytics pipeline. Bronze preserves bounded change images, Silver deduplicates and pseudonymizes activity, and Gold tables calculate user, tool, error, latency, watchlist, and research-save metrics.",
        [
            "pipelines/bronze_lakebase_changes.py and pipelines/silver_agent_activity.py implement the incremental foundation.",
            "Gold modules cover daily active researchers, tool latency, error rate, watchlist changes, and research saves.",
            "tools/phase6_cdf_acceptance.py verifies exact change counts, latency, aggregation, and cleanup.",
        ],
        [
            "A self-cleaning transaction produced 18 source history rows, 34 Bronze rows, 26 effective Silver rows, and exact results across five Gold metric families.",
            "Maximum reported source latency was 85 seconds. The submission claims Volume and Variety for the Big Data score rather than claiming sub-minute Velocity.",
        ],
    )

    add_category(
        doc,
        "Frontend and Core Workflow",
        "10 points",
        "The responsive Flask interface supports authenticated research, market and peer workflows, source inspection, watchlist actions, research memory, current signals, and usage analytics. Consequential writes have explicit controls and bounded feedback.",
        [
            "dashboard/templates/index.html and dashboard/static/app.js implement the complete interaction model.",
            "dashboard/auth.py enforces Google OIDC identity; dashboard/mcp_client.py creates a fresh signed assertion for each MCP session.",
            "Frontend and route contracts are covered by tests/test_dashboard_routes.py, tests/test_dashboard_ui_contract.py, and tests/test_render_frontend_routes.py.",
        ],
        [
            "The live application showed the expected signed-in user and loaded attributable research, an owned watchlist, saved notes and reports, latest signals, and Gold analytics.",
            "The interface surfaced an explicit stale-data status while preserving the most recent processed analytics, demonstrating a useful degraded state.",
        ],
    )

    add_category(
        doc,
        "Deployed Application",
        "5 points",
        "The frontend and FastMCP server run as separate Render services. Databricks retains the data and processing plane. The deployment contract includes health checks, reproducible builds, secret placeholders, OIDC configuration, and an operations runbook.",
        [
            "render.yaml defines both Python services and contains no Databricks data resources.",
            "docs/RENDER_DEPLOYMENT_RUNBOOK.md and docs/RENDER_OIDC_PREPARATION.md document setup, configuration, secret entry, validation, and recovery.",
            "docs/RENDER_DEPLOYMENT_PLAN.md records the boundary between Render and Databricks.",
        ],
        [
            "The frontend was accessible at the submitted URL under the verified Google account.",
            "Render showed signal-desk-mcp as Live, Blueprint managed, and linked to the deployed Git commit.",
            "Both services use Free instances; they must be pre-warmed because cold starts may add 50 seconds or more.",
        ],
    )

    add_category(
        doc,
        "Big Data Volume and Variety",
        "15 points",
        "The project demonstrates two rubric-defined Big Data characteristics with measurable downstream use. Volume comes from the distributed market dataset. Variety comes from processing structured market and XBRL data, semi-structured source payloads, and unstructured filing and news text.",
        [
            "The Spark market pipeline processes more than one million rows and publishes certified market history used by research workflows.",
            "Document text is normalized, section-aware, chunked, embedded, indexed, filtered, reranked, and returned with provenance.",
            "The accepted AI Search evaluation used 51 cases against a 393-row canonical serving corpus.",
        ],
        [
            "Volume: 1,255,489 unique Silver market rows after deterministic duplicate and malformed-record handling.",
            "Variety: 507 traceable filing and article chunks in the bounded research pipeline and 393 canonical search documents in the accepted index.",
            "Search quality: Recall at 5 of 1.0000, MRR of 0.9902, nDCG at 5 of 0.9928, zero provenance failures, and zero filter violations.",
        ],
    )


def add_live_screenshots(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Live Application Evidence", level=1)
    doc.add_paragraph(
        "The screenshots in this section were captured from the deployed application and Render dashboard on October 3 and 4, 2026, and the evidence package was revalidated on October 7, 2026. They contain no credentials, access tokens, connection strings, or private research bodies."
    )
    add_picture(
        doc,
        SCREENSHOTS / "app-research.png",
        6.85,
        "Authenticated Signal Desk research page showing the AAPL ticker performance workflow.",
        "Figure 2  Authenticated ticker performance workflow",
    )
    doc.add_page_break()
    add_picture(
        doc,
        SCREENSHOTS / "app-watchlist-signals.png",
        6.85,
        "Signal Desk page showing a three-ticker watchlist, latest signals, and user-owned research memory.",
        "Figure 3  User-owned watchlist signals notes and reports",
    )
    doc.add_paragraph(
        "The watchlist combines the latest locally served price with an as-of timestamp. The adjacent research memory confirms that notes and reports are retained for the authenticated user. Signals link to the underlying articles."
    )
    doc.add_page_break()
    add_picture(
        doc,
        SCREENSHOTS / "app-usage-analytics.png",
        6.15,
        "Signal Desk usage analytics showing daily active researchers, error rate, watchlist changes, research saves, and per-tool latency.",
        "Figure 4  Usage analytics from Databricks Gold tables",
    )
    doc.add_paragraph(
        "The screenshot shows current Gold metrics, including one daily active researcher, a 0.0 percent agent error rate over 15 invocations, 16 watchlist changes, 33 research saves, and per-tool P95 latency."
    )
    doc.add_page_break()
    add_picture(
        doc,
        SCREENSHOTS / "render-live-deployment.png",
        6.85,
        "Render deployment dashboard showing the Signal Desk MCP service as Live and Blueprint managed with successful deploys.",
        "Figure 5  Live Render MCP deployment",
    )
    doc.add_paragraph(
        "Render links the service to the GitHub repository and deployed commit. The visible Free plan warning is an operational constraint, not a functional failure; the runbook requires pre-warming before a demonstration."
    )


def add_quality_and_demo(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Verification and Release Evidence", level=1)
    doc.add_paragraph(
        "The submission combines current local verification with accepted workspace and deployed evidence. Local verification was rerun immediately before producing this package."
    )
    table = doc.add_table(rows=1, cols=3)
    for i, text in enumerate(("Check", "Result", "Meaning")):
        table.cell(0, i).text = text
    checks = [
        (
            "Automated tests",
            "281 passed",
            "Unit, contract, route, safety, pipeline-source, deployment, and acceptance behavior",
        ),
        ("Static analysis", "Ruff passed", "Repository-wide Python lint checks"),
        ("Supervisor evaluation", "10 of 10 passed", "Deployed action-taking agent and tool selection"),
        (
            "Retrieval evaluation",
            "51 cases accepted",
            "High recall, ranking quality, provenance, and metadata filtering",
        ),
        ("Credential scan", "Accepted release checkpoint", "No secrets in tracked source or sanitized evidence"),
        ("Live app inspection", "Verified", "Authenticated research, state, and analytics views loaded"),
    ]
    for values in checks:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(table, [1.45, 1.25, 3.95])

    doc.add_heading("Five Minute Demonstration Path", level=1)
    steps = [
        "Open the MCP URL first, then the frontend, and confirm the authenticated researcher.",
        "Run an AAPL filing and news question; inspect source type, date, score, and source URL.",
        "Run ticker performance and a two-company comparison on a shared window.",
        "Demonstrate a confirmed, disposable watchlist write and explain idempotent replay protection.",
        "Show research memory and the usage analytics metrics derived from Lakebase changes.",
        "Finish with the architecture diagram, the Volume and Variety counts, the Live Render deployment, and the 281-test result.",
    ]
    for step in steps:
        add_number(doc, step)

    doc.add_page_break()
    doc.add_heading("Operational Constraints", level=1)
    add_bullet(
        doc,
        "Render Free instances can sleep after inactivity. Pre-warm MCP first and frontend second before a live demo.",
    )
    add_bullet(
        doc,
        "Analytics may lag current UI actions. The interface labels stale data and shows the latest processed result.",
    )
    add_bullet(
        doc,
        "Massive plan entitlements can limit snapshots or fundamentals. The application labels unavailable data and falls back only through documented paths.",
    )
    add_bullet(
        doc,
        "The application supports research and evidence inspection, not trade execution or personalized investment advice.",
    )

    doc.add_heading("Submission Checklist", level=1)
    add_bullet(doc, "Upload Signal_Desk_Capstone_Submission.zip when submitting the complete evidence package.")
    add_bullet(
        doc, "Upload Signal_Desk_Capstone_Evidence.pdf instead when the grader prefers one directly viewable file."
    )
    add_bullet(doc, "Confirm the live application URL opens and the approved Google account can authenticate.")
    add_bullet(doc, "Pre-warm the MCP and frontend Render services before recording or presenting the demo.")
    add_bullet(
        doc, "Keep the five-minute demo script open and remove any disposable write created during the presentation."
    )


def add_index(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Evidence File Index", level=1)
    doc.add_paragraph(
        "The repository includes detailed phase status, architecture decisions, runbooks, tool references, and acceptance logic. The files below are the shortest path for a grader who wants to inspect the implementation behind a claim."
    )
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Evidence need"
    table.cell(0, 1).text = "Files"
    items = [
        ("Concise requirement map", "docs/release/REQUIREMENTS_TRACEABILITY.md"),
        ("Complete implementation history", "docs/IMPLEMENTATION_STATUS.md"),
        ("Deployment procedure", "docs/RENDER_DEPLOYMENT_RUNBOOK.md and render.yaml"),
        ("Architecture", "README.md and proposal/signal-desk-capstone-architecture.svg"),
        ("Data contracts", "shared/contracts/models.py and docs/release/DATA_DICTIONARY.md"),
        ("MCP tool behavior", "mcp_server/stock_research_mcp_server.py and docs/release/TOOL_API_REFERENCE.md"),
        ("Agent evaluation", "agent/evaluations.json and tools/phase7_agent_eval.py"),
        ("Spark definitions", "pipelines/ and resources/*.pipeline.yml"),
        ("Lakebase schema", "mcp_server/migrations/ and mcp_server/lakebase.py"),
        ("Analytics acceptance", "tools/phase6_cdf_acceptance.py and pipelines/gold_*"),
        ("Release checklist", "docs/release/DEMO_CHECKLIST.md"),
        ("Submission navigation", "submission/EVIDENCE_MANIFEST.md and submission/DEMO_SCRIPT.md"),
    ]
    for left, right in items:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    style_table(table, [2.0, 4.65])

    doc.add_heading("Submission Links", level=1)
    for label, url in (
        ("Signal Desk frontend", "https://signal-desk-frontend-s88i.onrender.com/"),
        ("Signal Desk MCP service", "https://signal-desk-mcp.onrender.com/"),
        ("Source repository", "https://github.com/srinivas-malyala/signal-desk-market-research"),
    ):
        p = doc.add_paragraph()
        add_link(p, label, url)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Prepared for the DataExpert Databricks AI Capstone rubric. ").bold = True
    p.add_run("All point claims are tied to observable behavior or a named repository artifact.")


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_overview(doc)
    add_score_map(doc)
    add_detailed_evidence(doc)
    add_live_screenshots(doc)
    add_quality_and_demo(doc)
    add_index(doc)
    doc.core_properties.title = "Signal Desk Capstone Evidence Report"
    doc.core_properties.subject = "Databricks AI Capstone grading evidence"
    doc.core_properties.author = "Srinivas Malyala"
    doc.core_properties.keywords = "Databricks, Spark, Lakebase, MCP, Render, capstone"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
